from docx import Document

doc=Document('高和萱程小奔20210408课评报告.docx')
table = doc.tables[0]
name = table.cell(1, 1).text

print(name)